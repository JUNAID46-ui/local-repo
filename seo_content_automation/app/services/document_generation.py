from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.models.article import FinalArticle
from app.models.business import BusinessProfile
from app.models.qa import QAResult

logger = logging.getLogger("seo_automation")


class DocumentGenerator:
    def generate(
        self,
        articles: list[FinalArticle],
        business: BusinessProfile,
        qa_results: list[QAResult],
        output_path: Path,
    ) -> None:
        doc = Document()

        style = doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

        self._add_cover(doc, business, len(articles))
        doc.add_page_break()
        self._add_business_info(doc, business)
        self._add_content_summary(doc, articles)
        doc.add_page_break()

        for i, (article, qa) in enumerate(zip(articles, qa_results), 1):
            self._add_article_section(doc, article, qa, i)
            if i < len(articles):
                doc.add_page_break()

        doc.add_page_break()
        self._add_run_summary(doc, articles, qa_results)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        logger.info(f"Document saved: {output_path}")

    def _add_cover(self, doc: Document, business: BusinessProfile, article_count: int) -> None:
        doc.add_paragraph()
        doc.add_paragraph()
        title = doc.add_heading(f"{business.business_name}", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle = doc.add_heading("SEO Content Package", level=1)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        from app.utils.dates import today_str
        info = doc.add_paragraph()
        info.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = info.add_run(f"Generated: {today_str()}\n")
        run.font.size = Pt(12)
        run = info.add_run(f"Articles: {article_count}\n")
        run.font.size = Pt(12)
        run = info.add_run(f"Location: {business.location}")
        run.font.size = Pt(12)

    def _add_business_info(self, doc: Document, business: BusinessProfile) -> None:
        doc.add_heading("Business Information", level=1)
        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"

        fields = [
            ("Business Name", business.business_name),
            ("Location", business.location),
            ("Industry", business.industry),
            ("Website", business.website),
            ("Services", ", ".join(business.services)),
            ("Service Areas", ", ".join(business.service_areas)),
            ("Target Audience", business.target_audience),
        ]
        for label, value in fields:
            if value:
                row = table.add_row()
                row.cells[0].text = label
                row.cells[1].text = value

    def _add_content_summary(self, doc: Document, articles: list[FinalArticle]) -> None:
        doc.add_heading("Content Summary", level=1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        headers = table.rows[0].cells
        headers[0].text = "#"
        headers[1].text = "Title"
        headers[2].text = "Primary Keyword"
        headers[3].text = "Words"
        headers[4].text = "KW Density"

        for i, article in enumerate(articles, 1):
            row = table.add_row()
            row.cells[0].text = str(i)
            row.cells[1].text = article.title
            row.cells[2].text = article.primary_keyword
            row.cells[3].text = str(article.word_count)
            row.cells[4].text = f"{article.keyword_density}%"

    def _add_article_section(
        self, doc: Document, article: FinalArticle, qa: QAResult, index: int,
    ) -> None:
        doc.add_heading(f"BLOG {index}", level=1)

        info_table = doc.add_table(rows=0, cols=2)
        info_table.style = "Light Grid Accent 1"
        meta = [
            ("Title", article.title),
            ("Primary Keyword", article.primary_keyword),
            ("Search Intent", article.search_intent),
            ("Service Category", article.service_category),
            ("Word Count", str(article.word_count)),
            ("Keyword Density", f"{article.keyword_density}%"),
            ("URL Slug", article.url_slug),
        ]
        for label, value in meta:
            row = info_table.add_row()
            row.cells[0].text = label
            row.cells[1].text = value

        doc.add_heading("Article", level=2)
        for line in article.body_text.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=2)
            else:
                doc.add_paragraph(line)

        doc.add_heading("Meta Description", level=2)
        p = doc.add_paragraph(article.meta_description)
        p.add_run(f"\n({len(article.meta_description)} characters)")

        if article.images:
            doc.add_heading("Image Ideas", level=2)
            for img in article.images:
                doc.add_heading(
                    f"Image {img.image_number}: {img.purpose.title()}", level=3,
                )
                img_table = doc.add_table(rows=0, cols=2)
                img_table.style = "Light Grid Accent 1"
                img_fields = [
                    ("Concept", img.concept),
                    ("Scene", img.scene_description),
                    ("Subject", img.subject),
                    ("Location Context", img.location_context),
                    ("Style", img.photography_style),
                    ("Filename", img.filename),
                    ("Alt Text", img.alt_text),
                    ("Placement", img.suggested_placement),
                ]
                for label, value in img_fields:
                    if value:
                        row = img_table.add_row()
                        row.cells[0].text = label
                        row.cells[1].text = value

        if article.internal_links:
            doc.add_heading("Internal Link Suggestions", level=2)
            link_table = doc.add_table(rows=1, cols=3)
            link_table.style = "Light Grid Accent 1"
            h = link_table.rows[0].cells
            h[0].text = "Anchor Text"
            h[1].text = "Target URL"
            h[2].text = "Reason"
            for link in article.internal_links:
                row = link_table.add_row()
                row.cells[0].text = link.anchor_text
                row.cells[1].text = link.target_url
                row.cells[2].text = link.reason

        if article.external_sources:
            doc.add_heading("External Sources", level=2)
            src_table = doc.add_table(rows=1, cols=3)
            src_table.style = "Light Grid Accent 1"
            h = src_table.rows[0].cells
            h[0].text = "Source"
            h[1].text = "URL"
            h[2].text = "Supports"
            for src in article.external_sources:
                row = src_table.add_row()
                row.cells[0].text = src.source_title
                row.cells[1].text = src.source_url
                row.cells[2].text = src.claim_supported

        if article.aeo_elements:
            doc.add_heading("AEO Opportunities", level=2)
            for elem in article.aeo_elements:
                doc.add_paragraph(elem, style="List Bullet")

        if article.geo_elements:
            doc.add_heading("GEO Opportunities", level=2)
            for elem in article.geo_elements:
                doc.add_paragraph(elem, style="List Bullet")

        doc.add_heading("QA Results", level=2)
        qa_table = doc.add_table(rows=0, cols=2)
        qa_table.style = "Light Grid Accent 1"
        qa_fields = [
            ("SEO QA", "Passed" if qa.seo.passed else "Issues Found"),
            ("AEO QA", "Passed" if qa.aeo.passed else "Issues Found"),
            ("GEO QA", "Passed" if qa.geo.passed else "Issues Found"),
            ("Fact Check", "Passed" if qa.fact_check.passed else "Issues Found"),
            ("Style Check", "Passed" if qa.style.passed else "Issues Found"),
            ("Overall", "PASSED" if qa.overall_passed else "NEEDS REVIEW"),
        ]
        for label, value in qa_fields:
            row = qa_table.add_row()
            row.cells[0].text = label
            row.cells[1].text = value

    def _add_run_summary(
        self, doc: Document, articles: list[FinalArticle], qa_results: list[QAResult],
    ) -> None:
        doc.add_heading("Run Summary", level=1)
        total = len(articles)
        passed = sum(1 for qa in qa_results if qa.overall_passed)
        total_words = sum(a.word_count for a in articles)
        avg_density = (
            sum(a.keyword_density for a in articles) / total if total > 0 else 0
        )

        summary_table = doc.add_table(rows=0, cols=2)
        summary_table.style = "Light Grid Accent 1"
        fields = [
            ("Total Articles", str(total)),
            ("QA Passed", f"{passed}/{total}"),
            ("Total Word Count", str(total_words)),
            ("Average Keyword Density", f"{avg_density:.2f}%"),
        ]
        for label, value in fields:
            row = summary_table.add_row()
            row.cells[0].text = label
            row.cells[1].text = value
